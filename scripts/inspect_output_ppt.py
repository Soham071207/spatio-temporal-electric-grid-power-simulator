from pptx import Presentation

prs = Presentation(r'EDI_MidSem_Presentation.pptx')

print(f'Total slides: {len(prs.slides)}')
print()

for slide_num, slide in enumerate(prs.slides):
    print(f'=== Slide {slide_num + 1} ===')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_text = []
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    all_text.append(t[:100])
            if all_text:
                print(f'  [{shape.name}] Texts: {all_text[:3]}...' if len(all_text) > 3 else f'  [{shape.name}] Texts: {all_text}')
        if shape.has_table:
            table = shape.table
            print(f'  [{shape.name}] TABLE: {len(list(table.rows))} rows x {len(table.columns)} cols')
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    ct = cell.text.strip()
                    if ct:
                        print(f'    Cell[{ri},{ci}]: "{ct[:60]}"')
    print()
