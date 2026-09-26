from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import json

prs = Presentation(r'PPT-TEMPLATE.pptx')

print(f'Slide width: {prs.slide_width}, height: {prs.slide_height}')
print(f'Slide width inches: {prs.slide_width / 914400}, height inches: {prs.slide_height / 914400}')
print(f'Number of slides: {len(prs.slides)}')
print()

print('=== SLIDE LAYOUTS ===')
for i, layout in enumerate(prs.slide_layouts):
    print(f'Layout {i}: {layout.name}')
    for ph in layout.placeholders:
        print(f'  Placeholder {ph.placeholder_format.idx}: {ph.name} ({ph.placeholder_format.type})')
print()

print('=== EXISTING SLIDES ===')
for slide_num, slide in enumerate(prs.slides):
    print(f'--- Slide {slide_num + 1} ---')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        print(f'  Shape: {shape.shape_type}, Name: {shape.name}, Left: {shape.left}, Top: {shape.top}, W: {shape.width}, H: {shape.height}')
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    runs_info = []
                    for run in para.runs:
                        font = run.font
                        runs_info.append({
                            'text': run.text,
                            'bold': font.bold,
                            'italic': font.italic,
                            'size': str(font.size) if font.size else None,
                            'color': str(font.color.rgb) if font.color and font.color.rgb else None,
                            'font_name': font.name,
                        })
                    print(f'    Text: "{text}"')
                    if runs_info:
                        print(f'    Runs: {json.dumps(runs_info, indent=6)}')
                    print(f'    Alignment: {para.alignment}')
        if hasattr(shape, 'image'):
            try:
                print(f'    [IMAGE] content_type={shape.image.content_type}, size={len(shape.image.blob)} bytes')
            except:
                pass
    print()
