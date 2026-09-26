from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import json

prs = Presentation(r'PPT-TEMPLATE.pptx')

print(f'Total slides: {len(prs.slides)}')
print()

for slide_num, slide in enumerate(prs.slides):
    print(f'=== Slide {slide_num + 1} ===')
    print(f'  Layout: {slide.slide_layout.name}')
    for shape in slide.shapes:
        if shape.has_text_frame:
            all_text = []
            for para in shape.text_frame.paragraphs:
                t = para.text.strip()
                if t:
                    all_text.append(t)
            if all_text:
                print(f'  [{shape.name}] Texts: {all_text}')
        if shape.has_table:
            table = shape.table
            print(f'  [{shape.name}] TABLE: {table.rows.__len__()} rows x {len(table.columns)} cols')
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    ct = cell.text.strip()
                    if ct:
                        print(f'    Cell[{ri},{ci}]: "{ct[:80]}..."' if len(ct)>80 else f'    Cell[{ri},{ci}]: "{ct}"')
    print()
